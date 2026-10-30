"""
Genera documentacion/entrega2/Entrega2_REDA.docx (Times New Roman 12, estilo de la Entrega 1).
El PDF se exporta después con Word (scripts/exportar_pdf.ps1) o LibreOffice.
Uso: python scripts/generar_documento_entrega2.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "documentacion", "entrega2", "Entrega2_REDA.docx")
AZUL = RGBColor(0x2E, 0x75, 0xB6)
REPO = "https://github.com/luismendoza009/REDA-Red-de-Evaluacion-y-Detecci-n-de-Riesgos-Urbanos"

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def titulo(texto, nivel=1):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(16 if nivel == 1 else 13)
    r.font.color.rgb = AZUL
    p.paragraph_format.space_before = Pt(12)
    return p


def parrafo(texto, cursiva=False, alinear=True):
    p = doc.add_paragraph()
    # **negrita** sencilla dentro del texto
    partes = texto.split("**")
    for i, parte in enumerate(partes):
        r = p.add_run(parte)
        r.bold = i % 2 == 1
        r.italic = cursiva
    if alinear:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def vineta(texto):
    p = parrafo("• " + texto)
    p.paragraph_format.left_indent = Cm(0.6)
    return p


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(encabezado, filas, anchos):
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezado):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        sombrear(c, "D9E7D3")
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            r = celdas[i].paragraphs[0].add_run(v)
            r.font.size = Pt(10)
            if v.startswith(("migraciones/", "sql/", "scripts/", "documentacion/", "diagramas/")):
                r.font.name = "Courier New"
                r.font.size = Pt(9)
    # Word solo respeta los anchos si el autoajuste está apagado y se fijan en columnas y celdas
    t.autofit = False
    for i, w in enumerate(anchos):
        t.columns[i].width = Cm(w)
    for fila in t.rows:
        for i, w in enumerate(anchos):
            fila.cells[i].width = Cm(w)
    doc.add_paragraph()


# ------------------------------------------------------------------ portada
for _ in range(6):
    doc.add_paragraph()
for texto, tam, negrita, cursiva in [
    ("UNIVERSIDAD POPULAR DEL CESAR", 14, True, False),
    ("Programa de Ingeniería de Sistemas", 12, False, False),
    ("", 12, False, False), ("", 12, False, False),
    ("REDA", 24, True, False),
    ("Red de Evaluación y Detección de Riesgos Urbanos", 13, False, True),
    ("", 12, False, False), ("", 12, False, False),
    ("Entrega 2 — Implementación", 14, True, False),
    ("", 12, False, False), ("", 12, False, False),
    ("Curso: Base de Datos I", 12, False, False),
    ("Integrantes: Rafael Alfonso Rivera Restrepo, Ivan Andrés Rodriguez Rodriguez, Luis Eliecer Mendoza Perpiñán, "
     "Carlos Eduardo Camelo Montaño", 12, False, False),
    ("", 12, False, False), ("", 12, False, False), ("", 12, False, False),
    ("Valledupar, Cesar — 2026", 12, False, False),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.size = Pt(tam)
    r.bold = negrita
    r.italic = cursiva
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ------------------------------------------------------------------ contenido
titulo("1. Qué contiene esta entrega")
parrafo("Esta entrega implementa en PostgreSQL 18 el diseño de la Entrega 1, ampliado por un cambio de enfoque del "
        "negocio: REDA pasa a ser un **sistema de documentación probatoria de fallas en infraestructura pública**. "
        "El ciudadano reporta condiciones (antecedentes) e incidentes (daños sufridos); el sistema vincula cada "
        "incidente con los reportes previos del mismo punto (\"la entidad ya sabía\") y, si el caso es viable, lo "
        "remite a un abogado. El esquema pasa de 10 a 25 tablas, y se conservan las 10 originales con sus políticas de borrado.")
parrafo("Los archivos de esta entrega se encuentran en la carpeta **base-de-datos/** del repositorio. La etiqueta **entrega-2** se creará después de revisar y aprobar la versión final. Las rutas de este documento son relativas a dicha carpeta:")
parrafo(REPO, alinear=False)

titulo("2. Dónde encontrar cada requisito")
tabla(["Requisito de la guía", "Ubicación (dentro de base-de-datos/)", "Resumen"], [
    ("1. Correcciones de la Entrega 1", "documentacion/entrega2/correcciones_entrega1.md", "Qué cambió en cada tabla y por qué"),
    ("2. DDL como migraciones numeradas", "migraciones/001_ … 017_*.sql", "17 migraciones: PK, FK con política, NOT NULL, UNIQUE, CHECK, índices, vistas y fechas DATE"),
    ("   Script de construcción", "scripts/migrar.ps1", "Base vacía → migraciones en orden (cada una en su transacción)"),
    ("   Evidencia de verificación", "documentacion/entrega2/evidencia_verificacion.txt", "Salida real: 17 migraciones, semilla y verificaciones completas"),
    ("3. Datos semilla", "sql/semilla/*.sql", "137 reportes de Valledupar, 18 incidentes, 10 casos; simulados y coherentes"),
    ("   Generador de la semilla", "scripts/generar_semilla.py", "Determinista; documentado en documentacion/datos_simulados.md"),
    ("4. Auditoría de normalización", "documentacion/entrega2/auditoria_normalizacion.md", "DF y forma normal de las 25 tablas, con 6 observaciones justificadas (E1–E6)"),
    ("5. Banco de consultas", "sql/consultas/banco_consultas.sql", "10 consultas comentadas con la pregunta de negocio"),
    ("   Verificación de consultas", "sql/consultas/verificacion_consultas.sql", "10 consultas verificadas por caminos independientes"),
    ("   Demostración de tres casos", "sql/pruebas/demostracion_tres_casos.sql", "Guion de sustentación: relaciones múltiples, ausencia de relación y rechazos"),
    ("   Evidencia de tres casos", "documentacion/entrega2/evidencia_tres_casos.txt", "Salida registrada de la demostración"),
    ("6. Índices con evidencia", "documentacion/entrega2/evidencia_indices.md", "2 experimentos EXPLAIN ANALYZE antes/después"),
    ("   Salida de los planes", "documentacion/entrega2/evidencia_indices.txt", "Planes reales de PostgreSQL 18.6"),
    ("7. Diccionario de datos", "documentacion/diccionario_datos.md", "Generado del esquema real con scripts/generar_diccionario.ps1"),
    ("Diagrama EER actualizado", "diagramas/reda_eer.drawio y .png", "25 tablas, 33 relaciones, notación pata de gallo"),
    ("Diagrama en notación de Chen", "diagramas/reda_eer_chen.drawio y .png", "Aporte de Rafael Rivera"),
    ("Decisiones técnicas", "documentacion/decisiones.md", "Decisiones D01–D34 y supuestos pendientes"),
    ("Pruebas de restricciones", "sql/pruebas/probar_restricciones.sql", "40 operaciones inválidas que la base debe rechazar"),
    ("Verificación desde cero", "scripts/verificar_desde_cero.ps1", "Base aparte: migraciones + semilla + todas las pruebas"),
], [4.6, 7.0, 4.4])

titulo("3. Cobertura del banco de consultas")
tabla(["Exigencia de la guía", "Consultas"], [
    ("2 consultas con JOIN de 3 o más tablas", "C01 (6 tablas), C02, C05, C08"),
    ("1 reunión externa \"los que no tienen\"", "C03 (incidentes validados sin caso), C04 (condiciones sin foto)"),
    ("1 agregación con GROUP BY y HAVING", "C06 (zonas con 3+ fallas abiertas), C10"),
    ("1 subconsulta o CTE", "C05 (CTE, antecedentes con Haversine), C07 (subconsulta), C08 (CTE)"),
    ("1 vista creada y consultada", "C09 sobre v_resumen_casos, C10 sobre v_mapa_publico"),
], [7, 9])

titulo("4. Cómo se verificó")
vineta("**Migraciones sobre base vacía:** la evidencia de verificación registra la aplicación de las 17 "
       "migraciones sin errores. scripts/verificar_desde_cero.ps1 permite repetir el proceso en una base aparte.")
vineta("**Integridad declarativa:** las 40 pruebas negativas de sql/pruebas/probar_restricciones.sql son rechazadas, cada una "
       "por la restricción exacta que se esperaba (la prueba compara el nombre de la restricción).")
vineta("**Coherencia de la semilla:** las 19 comprobaciones de sql/pruebas/verificar_semilla.sql devuelven 0 "
       "(estado actual = último cambio del historial, hecho antes que el reporte, antecedentes dentro de 100 m, etc.).")
vineta("**Consultas:** las 10 verificaciones cruzadas coinciden (por ejemplo, LEFT JOIN … IS NULL = NOT EXISTS = EXCEPT; "
       "el total facturado por abogado = el total cobrado por caso). C02 y C07 se comprobaron además a mano.")
vineta("**Tres casos de la rúbrica:** la demostración muestra registros relacionados con varios, registros sin relación "
       "y cuatro operaciones rechazadas por restricciones. La salida está en documentacion/entrega2/evidencia_tres_casos.txt.")
vineta("**Totales no inflados:** las vistas y consultas que combinan varias relaciones 1:N agregan en subconsultas "
       "antes de unir (ver v_resumen_casos y el comentario de C02).")
vineta("**Fechas:** las 26 fechas persistidas se declaran como DATE en la migración que crea cada tabla. "
       "Los antecedentes requieren un día previo y los cambios del historial se ordenan por id_historial (D34).")

titulo("5. Índices: resultados")
tabla(["Experimento", "Antes", "Después", "Mejora"], [
    ("Búsqueda de antecedentes (índice compuesto categoría, tipo, latitud)", "8,39 ms · 3 477 páginas", "0,54 ms · 574 páginas", "15× más rápido"),
    ("Fotos del expediente (índice de la FK evidencias.id_reporte)", "39,44 ms · Seq Scan de ≈ 400 000 filas", "0,069 ms · 8 páginas", "≈ 570× más rápido"),
], [6.4, 3.6, 3.6, 2.4])
parrafo("Las mediciones se hicieron con 200 000 reportes y 400 000 evidencias sintéticos cargados dentro de una "
        "transacción que termina en ROLLBACK, para que el volumen no contamine la base de la demo.")

titulo("6. Diagrama entidad-relación")
doc.add_picture(os.path.join(RAIZ, "diagramas", "reda_eer.png"), width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
parrafo("Figura 1. Modelo EER de REDA, Entrega 2. Archivo editable: diagramas/reda_eer.drawio.", cursiva=True)

titulo("7. Pendientes y riesgos declarados")
vineta("Reglas que las restricciones declarativas no alcanzan y que irán como triggers en la Entrega 3: que todo reporte "
       "INCIDENTE tenga su fila en incidentes, y que los pesos de cada versión de índice sumen 1.")
vineta("Funciones PL/pgSQL de cálculo de IRU e IVC y de búsqueda de antecedentes (Entrega 3).")
vineta("Los datos de la semilla son simulados (documentacion/datos_simulados.md). Los nombres de barrios y entidades están por "
       "verificar contra fuentes oficiales.")
vineta("Aspectos legales por validar con un abogado: modelo de cobro por caso remitido, plazo de caducidad de la acción de "
       "reparación directa y valor probatorio de fotos y geolocalización tomadas con una aplicación.")
vineta("Uso de inteligencia artificial: declarado en documentacion/declaracion_uso_ia.md.")

doc.save(SALIDA)
print("Generado", SALIDA)
